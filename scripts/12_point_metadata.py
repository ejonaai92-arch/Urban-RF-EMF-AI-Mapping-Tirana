"""Sampling design (Section 2.1, Fig. 1): location type (Cluster / Contour / Urban), field-sheet point number
and session start/end, linked to each dataset row through the point folder (Dt-<day>/P<nn>) of its spectrum file.
Input : data/field_sheets/Flete_Matje_21_dhjetori.xlsx (MKA), Flete_Matje_Zogu_Zi.xlsx (KT);
        data/spectra_source_files.csv (script 10)
Output: data/point_metadata.csv"""
import re
import openpyxl
import pandas as pd
from common import DATA, load_dataset

rows = []
for f, sq in (("Flete_Matje_21_dhjetori.xlsx", "MKA"), ("Flete_Matje_Zogu_Zi.xlsx", "KT")):
    wb = openpyxl.load_workbook(DATA / "field_sheets" / f, data_only=True)
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        ix = {k: hdr.index(k) for k in ("Point ID", "Start_Time", "End_Time", "Type")}
        ni = next(i for i, h in enumerate(hdr) if str(h).startswith("Narda_Saved_Index"))
        day = "Dt-" + ws.title
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] and str(r[0]).startswith("P"):
                rows.append(dict(Square=sq, Day_folder=day, Field_point=re.match(r"P\d+", str(r[ix["Point ID"]])).group(0),   # e.g. "P21* ©"
                                 Field_point_label=r[ix["Point ID"]], Narda_saved_index_sheet=str(r[ni]),
                                 Location_type=str(r[ix["Type"]]).strip().replace("Contural", "Contour"),
                                 Session_start=r[ix["Start_Time"]], Session_end=r[ix["End_Time"]]))
fs = pd.DataFrame(rows)
d = load_dataset()
src = pd.read_csv(DATA / "spectra_source_files.csv")
m = d[["Point_ID", "Square", "Date", "Time", "Weekday", "DayType", "X_UTM34N", "Y_UTM34N"]].merge(
    src[["Point_ID", "file", "Narda_index", "Instrument_time"]], on="Point_ID")
m["Day_folder"] = m.file.str.split("/").str[2]; m["Field_point"] = m.file.str.split("/").str[3].str.extract(r"^(P\d+)")[0]   # e.g. "P21(Qender_Zogu_i_Zi)"
m = m.merge(fs, on=["Square", "Day_folder", "Field_point"], how="left").drop(columns=["file"])
assert m.Location_type.notna().all() and len(m) == len(d)
print(m.groupby(["Square", "Location_type"]).size())
m.to_csv(DATA / "point_metadata.csv", index=False, float_format="%.2f")
